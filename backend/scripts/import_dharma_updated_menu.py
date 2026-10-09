"""
Load the final Sri Dharma Sastra menu onto the live store.

Reads "Menu Updated (2).xlsx" (sheet Reordered) and the photo folder
"all iamges" (including subfolders).

For every product row:
  - If that product is already on the store, keep it and sync the pack
    prices from the sheet (1 kg, 500 gram, 200 gram, 100 gram, per piece).
    Missing pack sizes are added. Packs the store already has are updated
    when the sheet price is different.
  - Category and subcategory are set to the storefront filter names, so
    the products page filter returns those items.
  - If it is not on the store, create it.
  - A photo from the image folder is attached when one matches.
  - No photo: the new product is saved as a draft, so it stays off
    https://kiterp.com/sri-dharma-sastra-home-foods/products

The second "Salted Sakinallu - Handmade" row is priced like the semi
version (Rs 269). The store cannot have two products with one name, so
that row is saved as "Salted Sakinallu - Semi Handmade".

From the backend folder:

    set DHARMA_LOGIN=vendor-phone-or-email
    set DHARMA_PASSWORD=secret
    py -3 scripts/import_dharma_updated_menu.py --apply

Clear every product tag without changing prices or variants:

    py -3 scripts/import_dharma_updated_menu.py --clear-tags
"""
from __future__ import annotations

import argparse
import json
import mimetypes
import os
import re
import sys
import uuid
from pathlib import Path
from urllib.error import HTTPError, URLError
from urllib.request import Request, urlopen

import openpyxl

API_BASE = os.environ.get("DHARMA_API", "https://kiterp.com/api/v1").rstrip("/")
VENDOR_SLUG = os.environ.get("DHARMA_VENDOR", "sri-dharma-sastra-home-foods")
DEFAULT_EXCEL = Path(r"c:\Users\hp\Downloads\Menu Updated (2).xlsx")
DEFAULT_IMAGES = Path(r"c:\Users\hp\Downloads\Menu - Dharma Sastra\all iamges")

MAX_IMAGE_BYTES = 5 * 1024 * 1024
STOCK_QTY = 100
IMAGE_EXTS = {".jpg", ".jpeg", ".png", ".webp", ".gif"}
EXT_RANK = {".jpg": 0, ".jpeg": 0, ".png": 1, ".webp": 2, ".gif": 3}

# Excel / import category text -> exact name on the products-page filter.
CATEGORY_MAP = {
    "pindi vantalu": "Pindi Vantalu",
    "milk + kova sweets": "Milk & Kova Based",
    "sweets": "Traditinal Sweets",
    "evening items": "Evening Snacks Hot",
    "snacks": "Snacks",
}

# Sheet name (normalized) -> existing store name (normalized).
ALIASES = {
    "karam sakinallu semi handmade": "karam sakinallu",
    "salted sakinallu handmade": "salted sakinallu handmade",
    "garjielu sugar rava coconut powder": "garijalu sugar rava coconut",
    "garjielu sugar kova coconut powder rava dry fruits": "kova gajakayalu",
    "garjielu bellam palli coconut powder": "garijalu peanut jaggery coconut",
    "garjielu bellam putnalu coconut powder": "garijalu roasted gram jaggery coconut",
    "garjielu bellam dry fruits coconut powder rava": "dry fruit garijelu jaggery dry fruits coconut rava",
    "small murukku": "micro murukku small",
    "murukku": "murukulu",
    "chekkara para": "chakkara parra",
    "rose cakes": "rose cookies",
    "mothichoor laddu": "motichoor laddu",
    "sunundalu bellam": "sunnundalu jaggery",
    "sunundalu chekkara": "sunnundalu sugar",
    "dry fruits laddu": "dry fruit laddu",
    "madata khaja": "madatha kaja",
    "azmeer kalakand": "ajmer kalakand",
    "diamond burfi": "bombay diamond burfi",
    "peda": "doodh peda",
    "double ka meeta": "double ka meetha",
    "bobbatlu kova": "bobbatlu khoya",
    "bobbatlu palli": "bobbatlu peanut",
    "bobbatlu pappu": "bobbatlu lentil",
    "samosa aloo": "aloo samosa",
    "samosa paneer": "samosa paneer",
    "samosa corn": "samosa corn",
    "samosa onion": "samosa onion",
    "pakodi onion": "onion pakodi",
    "pakodi palak": "palak pakoda",
    "jilebi sugar": "jalebi sugar syrup",
    "jilebi bellam": "jalebi jaggery syrup",
    "murukku ragulu": "raginalle murukku",
    "chakodi": "chakodi",
}

SIZES = (
    ("1kg", "1 kg", "kg", 1, 1.0),
    ("500g", "500 gram", "g", 500, 0.5),
    ("200g", "200 gram", "g", 200, 0.2),
    ("100g", "100 gram", "g", 100, 0.1),
    ("piece", "Per piece", "piece", 1, None),
)


def norm(value) -> str:
    text = str(value or "").lower().replace("—", " ").replace("–", " ")
    text = re.sub(r"\[[^\]]*\]", " ", text)
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return " ".join(text.split())


def clean_name(value) -> str:
    text = re.sub(r"\s*\[unclear\]\s*", " ", str(value or ""), flags=re.I)
    text = re.sub(r"\s+", " ", text).strip()
    text = re.sub(r"\s+\)", ")", text)
    text = re.sub(r"\(\s+", "(", text)
    return text


def store_category(raw: str) -> str:
    mapped = CATEGORY_MAP.get(raw.strip().lower())
    return mapped or raw.strip()


def price_num(value):
    if value is None or value == "":
        return None
    if isinstance(value, str):
        value = value.strip()
        if not value or value.lower() in {"no price", "-"}:
            return None
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    if number <= 0:
        return None
    return round(number, 2)


def covering_merge(ws, row: int, col: int):
    for merged in ws.merged_cells.ranges:
        if (merged.min_row <= row <= merged.max_row) and (merged.min_col <= col <= merged.max_col):
            return merged
    return None


def merged_value(ws, row: int, col: int):
    cell = ws.cell(row, col)
    if cell.value not in (None, ""):
        return cell.value
    merged = covering_merge(ws, row, col)
    if merged is None:
        return None
    return ws.cell(merged.min_row, merged.min_col).value


def is_banner(ws, row: int) -> bool:
    """A full-width merged title such as Snacks or Pickels, not a product name."""
    merged = covering_merge(ws, row, 3)
    return merged is not None and merged.min_col == 1 and merged.max_col >= 3


def load_sheet(path: Path) -> list[dict]:
    wb = openpyxl.load_workbook(path, data_only=True)
    if "Reordered" not in wb.sheetnames:
        raise SystemExit(f"{path} has no Reordered sheet")
    ws = wb["Reordered"]
    section = None
    rows = []
    seen_names: dict[str, int] = {}
    for excel_row in range(2, (ws.max_row or 1) + 1):
        sno = merged_value(ws, excel_row, 1)
        if is_banner(ws, excel_row):
            title = merged_value(ws, excel_row, 1)
            if isinstance(title, str) and title.strip():
                section = title.strip()
            continue
        sub_merge = covering_merge(ws, excel_row, 2)
        if sub_merge is not None and sub_merge.min_col == 1:
            sub = None
        else:
            sub = merged_value(ws, excel_row, 2)
        raw_name = ws.cell(excel_row, 3).value
        if raw_name is None or str(raw_name).strip() == "":
            if isinstance(sno, str) and sno.strip():
                section = sno.strip()
            continue
        name = clean_name(raw_name)
        key = norm(name)
        if key in seen_names and key == "salted sakinallu handmade":
            name = "Salted Sakinallu - Semi Handmade"
            key = norm(name)
        seen_names[key] = excel_row

        if section and "only by order" in section.lower():
            category = "Special order"
        elif isinstance(sub, str) and sub.strip():
            category = sub.strip()
        else:
            category = section or "Other"
        if category.lower() == "pickels":
            category = "Pickles"
        category = store_category(category)

        packs = []
        for idx, column in enumerate((4, 5, 6, 7, 8)):
            amount = price_num(ws.cell(excel_row, column).value)
            if amount is None:
                continue
            size_id, label, uom, qty, weight = SIZES[idx]
            packs.append({
                "id": size_id,
                "label": label,
                "uom": uom,
                "uom_quantity": qty,
                "weight_kg": weight,
                "price": amount,
            })
        if not packs:
            print(f"Row {excel_row}: {name} has no price, skipped", file=sys.stderr)
            continue
        rows.append({
            "excel_row": excel_row,
            "name": name,
            "key": key,
            "category": category,
            "packs": packs,
        })
    return rows


def list_images(folder: Path) -> list[Path]:
    files = []
    for path in folder.rglob("*"):
        if not path.is_file():
            continue
        if path.suffix.lower() not in IMAGE_EXTS:
            continue
        if path.stat().st_size > MAX_IMAGE_BYTES:
            print(f"Image over 5 MB, skipped: {path}", file=sys.stderr)
            continue
        files.append(path)
    return files


def _path_text(path: Path) -> str:
    return str(path).replace("\\", "/").lower()


def image_rule(key: str):
    rules = {
        "karam sakinallu handmade": lambda p: p.name.lower() in {"sakinallu.jpeg", "sakinallu.jpg"},
        "karam sakinallu semi handmade": lambda p: p.name.lower() in {"sakinallu.jpeg", "sakinallu.jpg"},
        "salted sakinallu handmade": lambda p: "salted sakinallu" in p.name.lower(),
        "salted sakinallu semi handmade": lambda p: "salted sakinallu" in p.name.lower(),
        "garjielu sugar rava coconut powder": lambda p: "sugar jargalu" in p.name.lower() or "sugar garjalu" in p.name.lower(),
        "garjielu sugar kova coconut powder rava dry fruits": lambda p: "kova gaja" in p.name.lower(),
        "garjielu bellam palli coconut powder": lambda p: "bellam garjalu" in p.name.lower(),
        "garjielu bellam dry fruits coconut powder rava": lambda p: "dry fruit" in p.name.lower() and "laddu" not in p.name.lower() and "pindi" in _path_text(p),
        "small murukku": lambda p: "chinna muruku" in p.name.lower(),
        "murukku": lambda p: p.name.lower() in {"murukulu.jpg", "murukulu.jpeg", "murukulu.png", "murukulu.webp"},
        "beyam murukku": lambda p: "beyam muruku" in p.name.lower(),
        "menappan murukku": lambda p: "minapa muruku" in p.name.lower(),
        "pappu chekkalu pachikaram": lambda p: "pappu karam chakralu" in p.name.lower(),
        "pappu chekkalu small chekkalu": lambda p: "small chekkalu" in p.name.lower(),
        "arisalu plain": lambda p: "plain arisa" in p.name.lower(),
        "arisalu nuvulu": lambda p: "nuvulu arisa" in p.name.lower(),
        "mixture": lambda p: "misture" in p.name.lower() and "garlic" not in p.name.lower(),
        "garlic mixture": lambda p: "garlic mixture" in p.name.lower(),
        "kara boondi": lambda p: "kara boondi" in p.name.lower(),
        "namak para": lambda p: "namak para" in p.name.lower(),
        "karam para": lambda p: "karam para" in p.name.lower(),
        "chekkara para": lambda p: "chekkara" in p.name.lower(),
        "sana ghati": lambda p: "sanna ghati" in p.name.lower(),
        "ghati": lambda p: "ghatiya" in p.name.lower(),
        "papadi": lambda p: p.name.lower().startswith("papadi"),
        "kara sev": lambda p: "kara sev" in p.name.lower(),
        "sanna sev": lambda p: "sanna sev" in p.name.lower(),
        "batani green pea": lambda p: "green pea" in p.name.lower(),
        "masala palli": lambda p: "masala palli" in p.name.lower(),
        "chana dal": lambda p: "chana dal" in p.name.lower(),
        "corn": lambda p: "makka" in p.name.lower(),
        "khatta meththa": lambda p: "khatta" in p.name.lower(),
        "paper attukulu": lambda p: "attukulu" in p.name.lower(),
        "bellam gavvalu": lambda p: "bellam gavvalu" in p.name.lower(),
        "sugar gavvalu": lambda p: "sugar gavvalu" in p.name.lower(),
        "karam gavvalu": lambda p: "guvvalu" in p.name.lower() or ("gavvalu" in p.name.lower() and "karam" in p.name.lower()),
        "salted chips": lambda p: "potato" in p.name.lower(),
        "hot chips": lambda p: "hot" in p.name.lower() and ("chip" in p.name.lower() or "chp" in p.name.lower()),
        "onion rings": lambda p: "/chips/" in _path_text(p) and "onion" in p.name.lower(),
        "ginger aloo fry": lambda p: "/chips/" in _path_text(p) and "ginger" in p.name.lower(),
        "jangari": lambda p: "jangari" in p.name.lower(),
        "mysore pak": lambda p: "mysore" in p.name.lower() and "milk" not in p.name.lower(),
        "rose cakes": lambda p: "rose cookie" in p.name.lower(),
        "sweet boondi": lambda p: "sweet boondi" in p.name.lower(),
        "boondi laddu": lambda p: "boondi laddu" in p.name.lower(),
        "mothichoor laddu": lambda p: "mothi" in p.name.lower(),
        "kova puri": lambda p: "kova puri" in p.name.lower(),
        "chandrakala": lambda p: "chandrakala" in p.name.lower(),
        "nuvvula laddu": lambda p: "laddu" in p.name.lower() and "nuvu" in p.name.lower() and "200" not in p.name.lower(),
        "besan laddu": lambda p: "besan laddu" in p.name.lower(),
        "rava laddu": lambda p: "laddu" in p.name.lower() and ("rava" in p.name.lower() or "raavva" in p.name.lower()),
        "sunundalu bellam": lambda p: "bellam sunnundalu" in p.name.lower(),
        "sunundalu chekkara": lambda p: "sugar sunundalu" in p.name.lower(),
        "sunundalu ragi": lambda p: "ragi laddu" in p.name.lower(),
        "dry fruits laddu": lambda p: "dry fruit" in p.name.lower() and "laddu" in p.name.lower(),
        "madata khaja": lambda p: "madatha" in p.name.lower(),
        "gottam khaja": lambda p: "gottam" in p.name.lower(),
        "malai khaja": lambda p: "malai khaja" in p.name.lower(),
        "chitti khaja": lambda p: "chitti khaja" in p.name.lower(),
        "kalakand": lambda p: "kalakand" in p.name.lower() and "ajmer" not in p.name.lower(),
        "azmeer kalakand": lambda p: "ajmer" in p.name.lower() and "kalakand" in p.name.lower(),
        "malai burfi": lambda p: "malai barfi" in p.name.lower() or "malai burfi" in p.name.lower(),
        "diamond burfi": lambda p: "dimound" in p.name.lower() or "diamond" in p.name.lower(),
        "badam barfi": lambda p: "badam barfi" in p.name.lower(),
        "peda": lambda p: "doodh peda" in p.name.lower(),
        "kala jamun": lambda p: "kala jamun" in p.name.lower(),
        "gulab jamun": lambda p: "gulab jamun" in p.name.lower(),
        "double ka meeta": lambda p: "double ka" in p.name.lower(),
        "coconut santhara": lambda p: "santra" in p.name.lower() or "santhara" in p.name.lower(),
        "milk mysore pak": lambda p: "milk mysore" in p.name.lower(),
        "idly karam": lambda p: "idli" in p.name.lower(),
        "sarva pindi pachi karam": lambda p: "pachi" in p.name.lower() and "sarva" in p.name.lower(),
        "sarva pindi yendu karam": lambda p: "karam sarva" in p.name.lower(),
        "bobbatlu kova": lambda p: "bobbat" in _path_text(p) and p.name.lower().startswith("ko"),
        "bobbatlu palli": lambda p: "palli bob" in p.name.lower(),
        "bobbatlu pappu": lambda p: "pappu bob" in p.name.lower(),
        "samosa aloo": lambda p: "aloo sam" in p.name.lower(),
        "samosa paneer": lambda p: "paneer" in p.name.lower(),
        "samosa corn": lambda p: "evening" in _path_text(p) and p.name.lower().startswith("corn"),
        "kachori kachori": lambda p: "kachori" in p.name.lower(),
        "kachori": lambda p: "kachori" in p.name.lower(),
        "pakodi onion": lambda p: "onion pakodi" in p.name.lower(),
        "pakodi palak": lambda p: "palak pakodi" in p.name.lower(),
        "jilebi sugar": lambda p: "sugar jel" in p.name.lower() or "sugar jal" in p.name.lower(),
        "jilebi bellam": lambda p: p.name.lower().startswith("jel") or p.name.lower().startswith("jal"),
        "chapati minimum 10": lambda p: "chapati" in p.name.lower(),
        "murukku ragulu": lambda p: "ragi muruku" in p.name.lower(),
    }
    return rules.get(key)


def match_image(product: dict, files: list[Path]) -> Path | None:
    rule = image_rule(product["key"])
    if rule is None:
        return None
    hits = [path for path in files if rule(path)]
    if not hits:
        return None
    hits.sort(key=lambda path: (EXT_RANK.get(path.suffix.lower(), 9), len(path.name)))
    return hits[0]


def content_type(path: Path) -> str:
    guessed, _ = mimetypes.guess_type(path.name)
    if guessed:
        return guessed
    return {
        ".jpg": "image/jpeg",
        ".jpeg": "image/jpeg",
        ".png": "image/png",
        ".webp": "image/webp",
        ".gif": "image/gif",
    }.get(path.suffix.lower(), "application/octet-stream")


def http_json(method: str, url: str, token: str | None = None, body=None, timeout: int = 120):
    data = None
    headers = {"X-Vendor-Slug": VENDOR_SLUG, "Accept": "application/json"}
    if token:
        headers["Authorization"] = f"Bearer {token}"
    if body is not None:
        data = json.dumps(body).encode("utf-8")
        headers["Content-Type"] = "application/json"
    request = Request(url, data=data, headers=headers, method=method)
    try:
        with urlopen(request, timeout=timeout) as response:
            raw = response.read().decode("utf-8")
            return response.status, json.loads(raw) if raw else {}
    except HTTPError as exc:
        detail = exc.read().decode("utf-8", errors="replace")
        raise RuntimeError(f"{method} {url} -> {exc.code}: {detail[:500]}") from exc
    except URLError as exc:
        raise RuntimeError(f"{method} {url} failed: {exc}") from exc


def encode_multipart(fields: dict, files: list[tuple[str, str, str, bytes]]) -> tuple[bytes, str]:
    boundary = f"----dharma{uuid.uuid4().hex}"
    chunks: list[bytes] = []
    for name, value in fields.items():
        chunks.append(
            f"--{boundary}\r\nContent-Disposition: form-data; name=\"{name}\"\r\n\r\n{value}\r\n".encode("utf-8")
        )
    for field, filename, mime, payload in files:
        chunks.append(
            (
                f"--{boundary}\r\n"
                f"Content-Disposition: form-data; name=\"{field}\"; filename=\"{filename}\"\r\n"
                f"Content-Type: {mime}\r\n\r\n"
            ).encode("utf-8")
        )
        chunks.append(payload)
        chunks.append(b"\r\n")
    chunks.append(f"--{boundary}--\r\n".encode("utf-8"))
    return b"".join(chunks), boundary


def http_multipart(url: str, token: str, fields: dict, files: list[tuple[str, str, str, bytes]], timeout: int = 180):
    payload, boundary = encode_multipart(fields, files)
    headers = {
        "Authorization": f"Bearer {token}",
        "X-Vendor-Slug": VENDOR_SLUG,
        "Content-Type": f"multipart/form-data; boundary={boundary}",
        "Accept": "application/json",
    }
    request = Request(url, data=payload, headers=headers, method="POST")
    try:
        with urlopen(request, timeout=timeout) as response:
            raw = response.read().decode("utf-8")
            return response.status, json.loads(raw) if raw else {}
    except HTTPError as exc:
        detail = exc.read().decode("utf-8", errors="replace")
        raise RuntimeError(f"POST {url} -> {exc.code}: {detail[:500]}") from exc


def login(email: str, password: str) -> str:
    _, body = http_json(
        "POST",
        f"{API_BASE}/auth/login",
        body={"login": email, "password": password, "vendor_slug": VENDOR_SLUG},
    )
    token = body.get("access_token")
    if not token:
        raise RuntimeError("Login did not return an access token")
    return token


def fetch_public_products() -> list[dict]:
    items = []
    page = 1
    while True:
        _, body = http_json("GET", f"{API_BASE}/catalog/products?page={page}&size=100")
        batch = body.get("items") or []
        items.extend(batch)
        pages = body.get("pages") or 1
        if page >= pages or not batch:
            break
        page += 1
    return items


def fetch_vendor_products(token: str) -> list[dict]:
    items = []
    page = 1
    while True:
        _, body = http_json(
            "GET",
            f"{API_BASE}/vendors/me/products?page={page}&size=200",
            token=token,
        )
        batch = body.get("items") or []
        items.extend(batch)
        pages = body.get("pages") or 1
        if page >= pages or not batch:
            break
        page += 1
    return items


def index_products(products: list[dict]) -> dict[str, dict]:
    indexed = {}
    for product in products:
        key = norm(product.get("name"))
        indexed.setdefault(key, product)
    return indexed


def find_existing(product: dict, indexed: dict[str, dict]) -> dict | None:
    if product["key"] in indexed:
        return indexed[product["key"]]
    alias = ALIASES.get(product["key"])
    if alias and alias in indexed:
        return indexed[alias]
    return None


def variant_size_key(variant: dict) -> str | None:
    name = norm(variant.get("name"))
    weight = variant.get("weight_kg")
    uom = norm(variant.get("uom"))
    qty = variant.get("uom_quantity")
    if weight is not None:
        try:
            w = float(weight)
        except (TypeError, ValueError):
            w = None
        if w is not None:
            if abs(w - 1) < 0.02:
                return "1kg"
            if abs(w - 0.5) < 0.02:
                return "500g"
            if abs(w - 0.2) < 0.02:
                return "200g"
            if abs(w - 0.1) < 0.02:
                return "100g"
    if qty is not None:
        try:
            q = float(qty)
        except (TypeError, ValueError):
            q = None
        if q is not None:
            if uom == "kg":
                if abs(q - 1) < 0.02:
                    return "1kg"
                if abs(q - 0.5) < 0.02:
                    return "500g"
            if uom in {"g", "gram", "grams"} or "gram" in uom:
                if abs(q - 1000) < 1:
                    return "1kg"
                if abs(q - 500) < 1:
                    return "500g"
                if abs(q - 200) < 1:
                    return "200g"
                if abs(q - 100) < 1:
                    return "100g"
    if "per piece" in name or name == "piece":
        return "piece"
    if "500" in name:
        return "500g"
    if "200" in name:
        return "200g"
    if "100" in name and "1 kg" not in name:
        return "100g"
    if name in {"1 kg", "1kg", "per kg"}:
        return "1kg"
    return None


def primary_pack(packs: list[dict]) -> dict:
    order = ["1kg", "500g", "200g", "100g", "piece"]
    by_id = {pack["id"]: pack for pack in packs}
    for size_id in order:
        if size_id in by_id:
            return by_id[size_id]
    return packs[0]


# Same option Sakinalu uses. The store card prints this as WEIGHT / 1 KG / 500 G.
WEIGHT_CHIPS = {
    "1kg": "1 kg",
    "500g": "500 g",
    "200g": "200 g",
    "100g": "100 g",
}


def product_uses_weight(packs: list[dict]) -> bool:
    return sum(1 for pack in packs if pack["id"] in WEIGHT_CHIPS) >= 2


def variant_payload(pack: dict, sku: str, chip: str | None = None, quantity: int | None = None) -> dict:
    body = {
        "name": pack["label"],
        "sku": sku,
        "uom": pack["uom"],
        "uom_quantity": pack["uom_quantity"],
        "price": pack["price"],
        "currency": "INR",
        "quantity": STOCK_QTY if quantity is None else quantity,
        "stock_status": "in_stock",
        "track_inventory": True,
        "weight_kg": pack["weight_kg"],
        "is_active": True,
    }
    if chip:
        body["attributes"] = {"Weight": chip}
    return body


def set_weight_option(payload: dict, chip: str | None) -> bool:
    """Write Weight like Sakinalu, and remove Size so the card does not ignore the pack."""
    attrs = dict(payload.get("attributes") or {})
    changed = False
    if "Size" in attrs:
        attrs.pop("Size", None)
        changed = True
    if chip and attrs.get("Weight") != chip:
        attrs["Weight"] = chip
        changed = True
    payload["attributes"] = attrs
    return changed


def kept_variant(existing: dict, pack: dict | None) -> dict:
    payload = {
        "id": existing.get("id"),
        "name": existing.get("name") or "Option",
        "sku": existing.get("sku"),
        "barcode": existing.get("barcode"),
        "uom": existing.get("uom") or "piece",
        "uom_quantity": existing.get("uom_quantity"),
        "price_type": existing.get("price_type") or "per_unit",
        "price": existing.get("price") or 0,
        "currency": existing.get("currency") or "INR",
        "quantity": existing.get("quantity") or 0,
        "low_stock_threshold": existing.get("low_stock_threshold") or 5,
        "stock_status": existing.get("stock_status") or "in_stock",
        "track_inventory": existing.get("track_inventory") if existing.get("track_inventory") is not None else True,
        "allow_backorders": existing.get("allow_backorders") or False,
        "weight_kg": existing.get("weight_kg"),
        "is_active": existing.get("is_active") if existing.get("is_active") is not None else True,
        "attributes": existing.get("attributes") or {},
    }
    if pack:
        payload["name"] = pack["label"]
        payload["uom"] = pack["uom"]
        payload["uom_quantity"] = pack["uom_quantity"]
        payload["price"] = pack["price"]
        payload["weight_kg"] = pack["weight_kg"]
    return payload


def product_body(product: dict, status: str) -> dict:
    primary = primary_pack(product["packs"])
    pack_text = ", ".join(f"{pack['label']} Rs {pack['price']:g}" for pack in product["packs"])
    body = {
        "name": product["name"],
        "description": f"{product['category']}. {pack_text}.",
        "short_description": product["category"],
        "category": product["category"],
        "subcategory": product["category"],
        "currency": "INR",
        "price": primary["price"],
        "uom": primary["uom"],
        "uom_quantity": primary["uom_quantity"],
        "weight_kg": primary["weight_kg"],
        "sku": f"DS-{product['excel_row']:03d}",
        "quantity": STOCK_QTY,
        "stock_status": "in_stock",
        "track_inventory": True,
        "status": status,
        "is_visible": status == "active",
    }
    # One variant is required for Add to Cart, even when the sheet has only 1 kg.
    use_weight = product_uses_weight(product["packs"])
    body["variants"] = [
        variant_payload(
            pack,
            f"DS-{product['excel_row']:03d}-{pack['id']}",
            WEIGHT_CHIPS.get(pack["id"]) if use_weight else None,
        )
        for pack in product["packs"]
    ]
    return body


def gram_changes(product: dict, existing: dict) -> tuple[list[dict], list[str]]:
    notes = []
    existing_variants = list(existing.get("variants") or [])
    sheet_by_size = {pack["id"]: pack for pack in product["packs"]}
    used_sizes = set()
    payload = []

    use_weight = product_uses_weight(product["packs"])

    if not existing_variants and product["packs"]:
        notes.append("add variant " + ", ".join(pack["label"] for pack in product["packs"]))
        payload = [
            variant_payload(
                pack,
                f"DS-{product['excel_row']:03d}-{pack['id']}",
                WEIGHT_CHIPS.get(pack["id"]) if use_weight else None,
            )
            for pack in product["packs"]
        ]
        return payload, notes

    for variant in existing_variants:
        size_id = variant_size_key(variant)
        pack = sheet_by_size.get(size_id) if size_id else None
        chip = WEIGHT_CHIPS.get(pack["id"]) if pack and use_weight else None
        if pack:
            used_sizes.add(size_id)
            old_price = round(float(variant.get("price") or 0), 2)
            if old_price != pack["price"]:
                notes.append(f"{pack['label']} Rs {old_price:g} -> Rs {pack['price']:g}")
                row = kept_variant(variant, pack)
            else:
                row = kept_variant(variant, None)
        else:
            row = kept_variant(variant, None)
        if set_weight_option(row, chip):
            notes.append(f"show {chip} on the store" if chip else "clear size option")
        payload.append(row)

    for pack in product["packs"]:
        if pack["id"] in used_sizes or not existing_variants:
            continue
        notes.append(f"add {pack['label']} Rs {pack['price']:g}")
        payload.append(
            variant_payload(
                pack,
                f"DS-{product['excel_row']:03d}-{pack['id']}",
                WEIGHT_CHIPS.get(pack["id"]) if use_weight else None,
            )
        )

    primary = primary_pack(product["packs"])
    old_product_price = round(float(existing.get("price") or 0), 2)
    if old_product_price != primary["price"]:
        notes.append(f"price Rs {old_product_price:g} -> Rs {primary['price']:g}")
    if (existing.get("category") or "") != product["category"]:
        notes.append(f"category -> {product['category']}")
    if (existing.get("uom") or "") != primary["uom"] or float(existing.get("uom_quantity") or 0) != float(primary["uom_quantity"]):
        notes.append(f"unit -> {primary['label']}")
    return payload, notes


def plan_rows(products, images, indexed):
    planned = []
    for product in products:
        photo = match_image(product, images)
        existing = find_existing(product, indexed) if indexed is not None else None
        planned.append({**product, "image": photo, "existing": existing})
    return planned


def print_plan(planned, have_variants: bool):
    create_active = create_draft = update = unchanged = 0
    print(f"{'Row':<5} {'Action':<16} {'Image':<5} {'Category':<22} Name")
    for item in planned:
        photo = "yes" if item["image"] else "no"
        existing = item["existing"]
        packs = ", ".join(f"{pack['label']} {pack['price']:g}" for pack in item["packs"])
        if existing is None:
            action = "CREATE active" if item["image"] else "CREATE draft"
            if item["image"]:
                create_active += 1
            else:
                create_draft += 1
            detail = packs
        else:
            if have_variants or "variants" in existing:
                _, notes = gram_changes(item, existing)
            else:
                notes = []
                old_price = round(float(existing.get("price") or 0), 2)
                new_price = primary_pack(item["packs"])["price"]
                if old_price != new_price:
                    notes.append(f"price Rs {old_price:g} -> Rs {new_price:g}")
                if len(item["packs"]) >= 2:
                    notes.append("sync pack sizes when logged in")
            needs_photo = item["image"] and not (existing.get("images") or [])
            if needs_photo:
                notes = list(notes) + ["attach photo"]
            if notes:
                action = "UPDATE"
                update += 1
                detail = "; ".join(notes)
            else:
                action = "UNCHANGED"
                unchanged += 1
                detail = existing.get("name") or ""
        print(f"{item['excel_row']:<5} {action:<16} {photo:<5} {item['category']:<22} {item['name']}", flush=True)
        if existing is not None and action == "UPDATE":
            print(f"{'':5} {'':16} {'':5} {'':22} existing: {existing.get('name')} | {detail}")
        elif existing is None:
            print(f"{'':5} {'':16} {'':5} {'':22} {detail}")
    print()
    print(
        f"Sheet products: {len(planned)} | "
        f"create active: {create_active} | create draft: {create_draft} | "
        f"update: {update} | unchanged: {unchanged}"
    )


def apply_plan(planned, token: str):
    created = updated = skipped = failed = 0
    for item in planned:
        try:
            existing = item["existing"]
            if existing is None:
                status = "active" if item["image"] else "draft"
                body = product_body(item, status)
                files = []
                if item["image"]:
                    path = item["image"]
                    files.append(("images", path.name, content_type(path), path.read_bytes()))
                http_multipart(
                    f"{API_BASE}/vendors/me/products",
                    token,
                    {"product_data": json.dumps(body), "primary_image_index": "0"},
                    files,
                )
                created += 1
                print(f"CREATED {status:6} {item['name']}", flush=True)
                continue

            payload, notes = gram_changes(item, existing)
            primary = primary_pack(item["packs"])
            changes = {}
            if round(float(existing.get("price") or 0), 2) != primary["price"]:
                changes["price"] = primary["price"]
            if (existing.get("uom") or "") != primary["uom"]:
                changes["uom"] = primary["uom"]
            if float(existing.get("uom_quantity") or 0) != float(primary["uom_quantity"]):
                changes["uom_quantity"] = primary["uom_quantity"]
            if existing.get("weight_kg") != primary["weight_kg"]:
                changes["weight_kg"] = primary["weight_kg"]
            if (existing.get("category") or "") != item["category"]:
                changes["category"] = item["category"]
                changes["subcategory"] = item["category"]
            if payload and any(
                note.startswith("add") or note.startswith("show ") or note.startswith("clear ") or "->" in note
                for note in notes
            ):
                changes["variants"] = payload
            needs_photo = bool(item["image"]) and not (existing.get("images") or [])
            if not changes and not needs_photo:
                skipped += 1
                print(f"SKIP     {item['name']}", flush=True)
                continue
            if changes:
                http_json(
                    "PUT",
                    f"{API_BASE}/vendors/me/products/{existing['id']}",
                    token=token,
                    body=changes,
                )
            if needs_photo:
                path = item["image"]
                http_multipart(
                    f"{API_BASE}/uploads/products/{existing['id']}/images",
                    token,
                    {},
                    [("file", path.name, content_type(path), path.read_bytes())],
                )
                if existing.get("status") == "draft":
                    http_json(
                        "PUT",
                        f"{API_BASE}/vendors/me/products/{existing['id']}",
                        token=token,
                        body={"status": "active", "is_visible": True},
                    )
            updated += 1
            print(f"UPDATED  {item['name']} | {'; '.join(notes) or 'photo'}", flush=True)
        except Exception as exc:
            failed += 1
            print(f"FAILED   {item['name']}: {exc}", file=sys.stderr)
    print()
    print(f"Created {created}, updated {updated}, unchanged {skipped}, failed {failed}")
    if failed:
        raise SystemExit(1)


def pack_from_product(product: dict) -> dict:
    """One sellable option from the product's own price and unit."""
    uom = product.get("uom") or "piece"
    try:
        qty = float(product.get("uom_quantity") or 1)
    except (TypeError, ValueError):
        qty = 1
    if uom == "kg" and abs(qty - 1) < 0.02:
        label = "1 kg"
    elif uom == "piece":
        label = "Per piece"
    elif uom in {"g", "gram", "grams"}:
        label = f"{qty:g} g"
    else:
        label = f"{qty:g} {uom}"
    try:
        stock = int(product.get("quantity") or 0)
    except (TypeError, ValueError):
        stock = 0
    return {
        "id": "base",
        "label": label,
        "uom": uom,
        "uom_quantity": qty,
        "weight_kg": product.get("weight_kg"),
        "price": round(float(product.get("price") or 0), 2),
        "quantity": stock if stock > 0 else STOCK_QTY,
    }


def product_tag_list(product: dict) -> list[str]:
    raw = product.get("tags")
    if not raw:
        return []
    if isinstance(raw, str):
        parts = raw.split(",")
    elif isinstance(raw, list):
        parts = raw
    else:
        return []
    return [str(part).strip() for part in parts if str(part).strip()]


def clear_product_tags(existing: list[dict], token: str | None, apply: bool):
    """Remove tag names from every product. Prices, variants, and status stay as they are."""
    tagged = [product for product in existing if product_tag_list(product)]
    print(f"Products with tags: {len(tagged)} of {len(existing)}")
    cleared = failed = 0
    for product in tagged:
        print(
            f"CLEAR TAG  {product.get('name')} | {', '.join(product_tag_list(product))}",
            flush=True,
        )
        if not apply:
            continue
        try:
            http_json(
                "PUT",
                f"{API_BASE}/vendors/me/products/{product['id']}",
                token=token,
                body={"tags": []},
            )
            cleared += 1
        except Exception as exc:
            failed += 1
            print(f"FAILED   {product.get('name')}: {exc}", file=sys.stderr)
    if apply:
        print(f"Tags cleared {cleared}, failed {failed}")
        if failed:
            raise SystemExit(1)


def products_missing_variant(existing: list[dict]) -> list[dict]:
    """Active products the sheet did not cover, still with nothing to add to the cart."""
    missing = []
    for product in existing:
        if product.get("status") != "active":
            continue
        active = [v for v in (product.get("variants") or []) if v.get("is_active") is not False]
        if not active:
            missing.append(product)
    return missing


def fill_products_without_variants(existing: list[dict], token: str | None, apply: bool):
    missing = products_missing_variant(existing)
    print()
    print(f"Active products with no variant: {len(missing)}")
    updated = failed = 0
    for product in missing:
        pack = pack_from_product(product)
        print(f"VARIANT  {product.get('name')} | {pack['label']} Rs {pack['price']:g}", flush=True)
        if not apply:
            continue
        try:
            payload = [kept_variant(variant, None) for variant in (product.get("variants") or [])]
            sku = f"{product.get('sku') or 'DS'}-1"
            payload.append(variant_payload(pack, sku, quantity=pack["quantity"]))
            http_json(
                "PUT",
                f"{API_BASE}/vendors/me/products/{product['id']}",
                token=token,
                body={"variants": payload},
            )
            updated += 1
        except Exception as exc:
            failed += 1
            print(f"FAILED   {product.get('name')}: {exc}", file=sys.stderr)
    if apply:
        print(f"Variants added {updated}, failed {failed}")
        if failed:
            raise SystemExit(1)
    return missing


def main():
    parser = argparse.ArgumentParser(description="Import the final Dharma Sastra menu")
    parser.add_argument("--excel", type=Path, default=DEFAULT_EXCEL)
    parser.add_argument("--images", type=Path, default=DEFAULT_IMAGES)
    parser.add_argument("--offline", action="store_true", help="Do not call the live catalog.")
    parser.add_argument("--apply", action="store_true", help="Create and update products. Default is a dry run.")
    parser.add_argument(
        "--clear-tags",
        action="store_true",
        help="Remove tag names from every product. Does not change prices or variants.",
    )
    parser.add_argument("--login", default=os.environ.get("DHARMA_LOGIN", ""))
    parser.add_argument("--password", default=os.environ.get("DHARMA_PASSWORD", ""))
    args = parser.parse_args()

    if args.clear_tags:
        if not args.login or not args.password:
            raise SystemExit("Set DHARMA_LOGIN and DHARMA_PASSWORD before --clear-tags")
        token = login(args.login, args.password)
        existing = fetch_vendor_products(token)
        print(f"Vendor catalog: {len(existing)} products (includes drafts)")
        clear_product_tags(existing, token, apply=True)
        return

    if not args.excel.is_file():
        raise SystemExit(f"Excel file not found: {args.excel}")
    if not args.images.is_dir():
        raise SystemExit(f"Image folder not found: {args.images}")

    products = load_sheet(args.excel)
    images = list_images(args.images)
    token = None
    have_variants = False
    if args.offline:
        existing = []
        print("Offline check: sheet and image folder only.")
    elif args.login and args.password:
        token = login(args.login, args.password)
        existing = fetch_vendor_products(token)
        have_variants = True
        print(f"Vendor catalog: {len(existing)} products (includes drafts)")
    else:
        if args.apply:
            raise SystemExit("Set DHARMA_LOGIN and DHARMA_PASSWORD before --apply")
        try:
            existing = fetch_public_products()
            print(f"Public catalog: {len(existing)} products. Drafts are not in this list.")
        except RuntimeError as exc:
            existing = []
            print(f"Could not read the live catalog ({exc}).")
    planned = plan_rows(products, images, index_products(existing))
    print_plan(planned, have_variants)
    if not args.apply:
        if have_variants:
            fill_products_without_variants(existing, None, apply=False)
        print("Dry run only. Re-run with --apply to write these products.")
        return
    if not token:
        raise SystemExit("Set DHARMA_LOGIN and DHARMA_PASSWORD before --apply")
    apply_plan(planned, token)
    fill_products_without_variants(fetch_vendor_products(token), token, apply=True)


if __name__ == "__main__":
    main()
